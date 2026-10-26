"""US Treasury Office of Tax Policy tax expenditure estimates (published with the President's Budget).

The latest edition is discovered on the Treasury landing page (Tax-Expenditures-FY{YYYY}.xlsx + .pdf).
Receipt (revenue-loss) effects per provision, US$ million, fiscal years (Oct-Sep), current year plus a
ten-year window, from workbook Table 1; the corporate / individual split from Tables 2a and 2b; outlay
effects of refundable credits from Table 5 (kept separate, never added to the receipt effect). Budget
function from the Table 1 headings; descriptions and the document date from the PDF, whose Table 1 is
checked value by value against the workbook. Treasury publishes no grand total and cautions against
adding provisions (interactions), so the official aggregate is null; quantified sums are labelled derived.
US Government work: public domain (17 U.S.C. 105).
"""
import collections
import difflib
import hashlib
import io
import re
from datetime import datetime

import openpyxl

from .text import long_date

LANDING = 'https://home.treasury.gov/policy-issues/tax-policy/tax-expenditures'
HOST = 'https://home.treasury.gov'
PUBLISHER = 'US Department of the Treasury, Office of Tax Policy (Office of Tax Analysis)'
LICENCE = 'Public domain (US Government work, 17 U.S.C. 105)'
ATTRIBUTION = 'Source: US Department of the Treasury, Office of Tax Analysis, Tax Expenditures.'
SHEETS = {'total': 'Table 1', 'corporate': 'Table 2a', 'individual': 'Table 2b', 'outlays': 'Table 5'}
TITLES = {'total': r'ESTIMATES OF TOTAL INCOME TAX EXPENDITURES', 'corporate': r'ESTIMATES OF TOTAL CORPORATE INCOME TAX EXPENDITURES',
          'individual': r'ESTIMATES OF TOTAL INDIVIDUAL INCOME TAX EXPENDITURES', 'outlays': r'ESTIMATES OF OUTLAY TAX EXPENDITURES'}
HEADS = {'CI': 'Individual and corporate income tax', 'I': 'Individual income tax', 'C': 'Corporate income tax', '': 'Individual and corporate income tax'}
RETIREMENT = 'Net exclusion of pension contributions and earnings'
MAX_DERIVED_SPLITS = 3  # Table 2a/2b rows that may be missing (split derived as Table 1 minus the other table).
STATUSES = ('estimate', 'forecast')
MAX_TITLE_MISMATCHES = 8
# Group subheadings ('Net exclusion of pension contributions and earnings:') cover only the next few provisions, unlike
# section subheadings ('Housing:'), which run to the next heading. Treasury marks the group's end only by indentation,
# which the workbook and PDF apply inconsistently; the first provision named as a provision in its own right closes it.
GROUP_HEADING = re.compile(r'(?:Net )?exclusion of\b', re.I)
GROUP_END = re.compile(r'(?:Additional|Temporary|Credit|Deduction|Deductibility|Deferral|Exclusion|Expensing|Exemption)\b', re.I)
MONTHS = '(?:January|February|March|April|May|June|July|August|September|October|November|December)'


def fy(year):
    """US fiscal year 2026 (Oct 2025-Sep 2026) -> '2025-26' (start calendar year, project convention)."""
    year = int(year)
    return f'{year - 1}-{str(year)[2:]}'


def links(html):
    """Landing page -> (fiscal-year edition, xlsx URL, pdf URL) of the latest edition."""
    if not re.search(r'<title>\s*Tax Expenditures\b', html):
        raise ValueError('Treasury tax expenditures landing page changed')
    found = {}
    for path, year, ext in re.findall(r'href=["\']?((?:https://home\.treasury\.gov)?/system/files/131/Tax-Expenditures-(?:FY)?(\d{4})(?:-update)?\.(xlsx|pdf))', html):
        found.setdefault(int(year), {}).setdefault(ext, set()).add(path if path.startswith('https://') else HOST + path)
    if not found:
        raise ValueError('No Treasury tax expenditure files linked from the landing page')
    edition = max(found)
    files = found[edition]
    if len(files.get('xlsx', ())) != 1 or len(files.get('pdf', ())) != 1:
        raise ValueError(f'Treasury FY{edition} tax expenditures: expected exactly one workbook and one PDF link')
    return edition, next(iter(files['xlsx'])), next(iter(files['pdf']))


def clean_name(text):
    text = re.sub(r'(\s+\d/)+\s*$', '', str(text).strip())
    return re.sub(r'\s+\.$', '', re.sub(r'\s+', ' ', text)).strip()


def key(name):
    return re.sub(r'[^a-z0-9]', '', clean_name(name).lower())


def similar(a, b):
    return difflib.SequenceMatcher(None, key(a), key(b)).ratio()


def parse_table(ws, kind, edition):
    """Workbook table -> (years, rows, footnotes). Rows keep number, name, function, subheading, values."""
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    title = next((str(r[0]) for r in rows if r and isinstance(r[0], str) and 'ESTIMATES OF' in r[0]), '')
    if not re.search(TITLES[kind], title) or not any(r and r[0] == '(in millions of dollars)' for r in rows):
        raise ValueError(f'Treasury {SHEETS[kind]} title or units changed: {title!r}')
    header = next((i for i, r in enumerate(rows) if len(r) > 2 and isinstance(r[2], int) and 1990 < r[2] < 2100), None)
    if header is None:
        raise ValueError(f'Treasury {SHEETS[kind]} year header missing')
    years = []
    for v in rows[header][2:]:
        if not isinstance(v, int):
            break
        years.append(v)
    window = rows[header][2 + len(years)] if len(rows[header]) > 2 + len(years) else None
    if years != list(range(edition - 2, edition + 9)) or window != f'{years[1]}-{years[-1]}':
        raise ValueError(f'Treasury {SHEETS[kind]} years changed: {years} {window!r}')
    width = len(years) + 1
    out, notes, function, section, group, addendum, last_note = [], {}, None, None, None, False, None
    for r in rows[header + 1:]:
        r = (r + [None] * (2 + width))[:2 + width]
        a, b, values = r[0], r[1], r[2:]
        if a is None and b is None and all(v is None for v in values):
            continue
        if isinstance(a, str) and re.fullmatch(r'\d+/', a.strip()):
            last_note = a.strip()
            notes[last_note] = str(b or '').strip()
            continue
        if last_note and a is None and isinstance(b, str) and all(v is None for v in values):
            notes[last_note] = (notes[last_note] + ' ' + b.replace('\xa0', ' ').strip()).strip()
            continue
        if isinstance(a, str) and all(v is None for v in values) and b is None:
            heading = a.strip().rstrip(':').strip()
            addendum = heading.lower().startswith('addendum')
            function, section, group = heading, None, None
            continue
        if addendum:
            continue  # Addendum (aid to State and local governments) repeats items already listed.
        if a is None and isinstance(b, str) and b.strip().endswith(':') and all(v in (None, 0) for v in values):
            heading = b.strip().rstrip(':').strip()  # Tables 2a/2b fill some subheading rows with zeros.
            if GROUP_HEADING.match(heading):
                group = heading
            else:
                section, group = heading, None
            continue
        if isinstance(a, int) and isinstance(b, str) and function:
            nums = values[:width]
            if kind == 'outlays':  # Blank outlay cells after a credit ends are nil.
                nums = [0 if v is None else v for v in nums]
            if any(not isinstance(v, (int, float)) or isinstance(v, bool) for v in nums):
                raise ValueError(f'Treasury {SHEETS[kind]} item {a}: non-numeric values {nums}')
            if any(v is not None for v in r[2 + width:]):
                raise ValueError(f'Treasury {SHEETS[kind]} item {a}: unexpected extra columns')
            if nums[-1] != sum(nums[1:-1]):
                raise ValueError(f'Treasury {SHEETS[kind]} item {a}: ten-year total does not add up')
            marks = re.findall(r'(\d+/)', re.search(r'((?:\s+\d/)*)\s*$', b)[1])
            if group and GROUP_END.match(clean_name(b)):
                group = None
            out.append({'number': a, 'rawName': b, 'name': clean_name(b), 'function': function, 'subheading': group or section,
                        'values': [float(v) for v in nums[:-1]], 'footnotes': marks})
            continue
        raise ValueError(f'Unexpected row in Treasury {SHEETS[kind]}: {r[:4]}')
    return [fy(y) for y in years], out, notes


def parse_workbook(raw, edition):
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    tables = {}
    for kind, prefix in SHEETS.items():
        sheet = [s for s in wb.sheetnames if s.startswith(prefix + ' ')]
        if len(sheet) != 1:
            raise ValueError(f'Treasury workbook sheet missing or ambiguous: {prefix}')
        tables[kind] = parse_table(wb[sheet[0]], kind, edition)
    if len({tuple(t[0]) for t in tables.values()}) != 1:
        raise ValueError('Treasury workbook tables cover different years')
    return tables


def pdf_text(raw):
    """PDF bytes -> list of page texts (pypdf)."""
    import pypdf
    return [p.extract_text() or '' for p in pypdf.PdfReader(io.BytesIO(raw)).pages]


def tidy(text):
    text = re.sub(r'\s+', ' ', text).strip()
    text = re.sub(r'([A-Za-z]) -([a-z])', r'\1-\2', text)
    text = re.sub(r'([A-Za-z])- ([a-z])', r'\1-\2', text)
    return re.sub(r'\s+([.,;:)])', r'\1', text)


def vocabulary(pages):
    """Lower-case word counts over the whole PDF text: the evidence for rejoining split words."""
    return collections.Counter(w.lower() for p in pages for w in re.findall(r'[A-Za-z]+', p))


def rejoin(text, words):
    """Repair words that PDF text extraction split with a stray space ('qu alifying', 'militar y', 't o').

    Two neighboring pieces are joined only when the joined word occurs elsewhere in the document and one piece
    is evidently a fragment: a stray single letter (not 'a' or 'I'), or a piece that is rare in the document. A common
    word followed by a long rare word ('in tangible') is left alone, as is anything before a capitalized word.
    """
    count = lambda w: words[w.lower()]
    rare = lambda w: count(w) <= 3
    letter = lambda w: len(w) == 1 and w not in 'aAI'

    def join(a, b):
        if (b[0].isupper() and not a.isupper()) or count(a + b) < 1:
            return False
        if a in ('a', 'A', 'I'):
            return rare(b)
        if b in ('a', 'I'):
            return rare(a)
        if letter(a) or letter(b):
            return True
        if not (rare(a) or rare(b)):
            return False
        return not ((count(a) > 3 and len(b) >= 5) or (count(b) > 3 and len(a) >= 5 and len(b) >= 3))

    def pieces(left, right):
        h, t = re.fullmatch(r'[(\["\'“]*([A-Za-z]+)', left), re.match(r'([A-Za-z]+)', right)
        return (h[1], t[1]) if h and t else None

    tokens, out = text.split(' '), []
    for i, token in enumerate(tokens):
        pair = pieces(out[-1], token) if out else None
        if pair and join(*pair):
            ahead = pieces(token, tokens[i + 1]) if i + 1 < len(tokens) else None
            # 'in t heir': a stray letter belongs to the following piece when that join is better attested.
            if not (ahead and len(pair[1]) == 1 and join(*ahead) and count(''.join(ahead)) > count(''.join(pair))):
                out[-1] += token
                continue
        out.append(token)
    return ' '.join(out)


def parse_pdf(pages, edition):
    """PDF pages -> (document date, {number: (title, description)}, {number: Table 1 values})."""
    if not re.search(rf'TAX EXPENDITURES\s+FISCAL YEAR\s+{edition}\b', pages[0] if pages else ''):
        raise ValueError(f'Treasury PDF is not the FY{edition} tax expenditures document')
    dates = set(re.findall(rf'Office of Tax Analysis\s*\n\s*({MONTHS} \d{{1,2}}, \d{{4}})', '\n'.join(pages)))
    if len(dates) != 1:
        raise ValueError(f'Treasury PDF document date missing or inconsistent: {sorted(dates)}')
    published = datetime.strptime(dates.pop(), '%B %d, %Y').date().isoformat()
    table_pages = [i for i, p in enumerate(pages) if 'Total from corporations and individuals' in p or 'ESTIMATES OF TOTAL INCOME TAX EXPENDITURES' in p]
    start = next((i for i, p in enumerate(pages) if 'Descriptions of Income Tax Provisions' in p), None)
    if start is None or not table_pages or table_pages[0] <= start:
        raise ValueError('Treasury PDF layout changed (descriptions or Table 1 not found)')
    lines = []
    for i in range(start, table_pages[0]):
        text = pages[i]
        if i == start:
            text = text.split('Descriptions of Income Tax Provisions', 1)[1]
        for line in text.split('\n'):
            s = line.strip()
            if re.match(r'\d{1,2}\s{2,}\S', line.lstrip()) and not re.match(r'\d+\.', s):
                break  # Footnotes run to the foot of the page.
            if not s or re.fullmatch(r'\d{1,2}', s) or s.startswith('U.S. Department of the Treasury') or s == 'Office of Tax Analysis' or re.fullmatch(rf'{MONTHS} \d{{1,2}}, \d{{4}}', s):
                continue
            lines.append(s)
    blocks, expected, skipping = {}, 1, False
    heading = re.compile(r'[A-Z][A-Za-z]*,?(?: (?:and|of|the|[A-Z][A-Za-z]*,?))*')
    for j, s in enumerate(lines):
        if s.startswith(f'{expected}. ') or s.startswith(f'{expected}.—'):
            blocks[expected] = [s[len(str(expected)) + 1:]]
            expected += 1
            skipping = False
        elif blocks and not skipping:
            if len(s) < 70 and heading.fullmatch(s) and re.search(r'[.)]$', lines[j - 1]):
                skipping = True  # Budget-function heading and its introductory paragraph, up to the next provision.
                continue
            blocks[expected - 1].append(s)
    descriptions, words = {}, vocabulary(pages)
    for n, parts in blocks.items():
        text = rejoin(tidy(' '.join(parts)), words)
        if '—' not in text:
            raise ValueError(f'Treasury PDF description {n} has no title separator')
        title, body = text.split('—', 1)
        descriptions[n] = (title.strip().rstrip('.').strip(), body.strip())
    table = {}
    row = re.compile(r'^(\d{1,3}) (.+?) ((?:-?[\d,]+ ){11}-?[\d,]+)\s*$')
    for i in table_pages:
        for line in pages[i].split('\n'):
            m = row.match(line.strip())
            if m:
                n = int(m[1])
                if n in table:
                    raise ValueError(f'Duplicate PDF Table 1 row {n}')
                table[n] = (clean_name(m[2]), [float(v.replace(',', '')) for v in m[3].split()][:-1])
    return published, descriptions, table


def parse_reliefs(xlsx, pages, edition, xlsx_url, pdf_url, retrieved=None):
    tables = parse_workbook(xlsx, edition)
    years, total, notes = tables['total']
    numbers = [r['number'] for r in total]
    if numbers != list(range(1, len(total) + 1)):
        raise ValueError('Treasury Table 1 provision numbers are not 1..N in order')
    published, descriptions, pdf_table = parse_pdf(pages, edition)
    first_forecast = fy(int(published[:4]) + (1 if int(published[5:7]) >= 10 else 0))
    if first_forecast not in years or years.index(first_forecast) != 1:
        raise ValueError(f'Treasury edition dated {published} does not start one year before the current fiscal year')
    if set(descriptions) != set(numbers) or set(pdf_table) != set(numbers):
        raise ValueError(f'Treasury PDF provisions differ from the workbook: descriptions {len(descriptions)}, table {len(pdf_table)}, workbook {len(numbers)}')
    issues = []
    split = {}
    for kind in ('corporate', 'individual'):
        rows = tables[kind][1]
        by_key = {}
        for r in rows:
            if key(r['name']) in by_key:
                raise ValueError(f'Duplicate provision name in Treasury {SHEETS[kind]}: {r["name"]}')
            by_key[key(r['name'])] = r
        matched = set()
        for t in total:
            r = by_key.get(key(t['name']))
            if r is None:
                split[(kind, t['number'])] = None
                continue
            matched.add(key(r['name']))
            if r['number'] != t['number']:
                issues.append(f'{SHEETS[kind]} numbers "{t["name"]}" as {r["number"]} (Table 1: {t["number"]}); matched by name.')
            split[(kind, t['number'])] = r['values']
        extra = set(by_key) - matched
        if extra:
            raise ValueError(f'Treasury {SHEETS[kind]} rows not in Table 1: {sorted(extra)[:5]}')
    derived = []
    for t in total:
        corp, ind = split[('corporate', t['number'])], split[('individual', t['number'])]
        if corp is None and ind is None:
            raise ValueError(f'Treasury provision {t["number"]} missing from both Tables 2a and 2b')
        if corp is None or ind is None:
            derived.append(t['number'])
            kind = 'corporate' if corp is None else 'individual'
            issues.append(f'"{t["name"]}" (provision {t["number"]}) is missing from {SHEETS[kind]}; its {kind} portion is derived as Table 1 minus the other table.')
            corp = corp if corp is not None else [a - b for a, b in zip(t['values'], ind)]
            ind = ind if ind is not None else [a - b for a, b in zip(t['values'], corp)]
        elif any(a + b != c for a, b, c in zip(corp, ind, t['values'])):
            raise ValueError(f'Treasury provision {t["number"]}: corporate + individual != total')
        t['corporate'], t['individual'] = corp, ind
    if len(derived) > MAX_DERIVED_SPLITS:
        raise ValueError(f'Too many provisions missing from Tables 2a/2b: {derived}')
    for t in total:
        name, values = pdf_table[t['number']]
        if values != t['values']:
            raise ValueError(f'Treasury provision {t["number"]}: workbook {t["values"]} disagrees with PDF Table 1 {values}')
        if similar(name, t['name']) < 0.9:
            raise ValueError(f'Treasury provision {t["number"]}: workbook name {t["name"]!r} disagrees with PDF {name!r}')
    # Description titles are sometimes worded differently (e.g. 'Deduction for qualified tips' for 'Exclusion of tip income');
    # a sequence drift would mismatch nearly every title, so only a few low-similarity titles are tolerated.
    loose = [t['number'] for t in total if similar(descriptions[t['number']][0], t['name']) < 0.6]
    if len(loose) > MAX_TITLE_MISMATCHES:
        raise ValueError(f'Treasury PDF description titles do not match the workbook for provisions {loose}')
    outlays = {}
    for r in tables['outlays'][1]:
        t = total[r['number'] - 1] if 1 <= r['number'] <= len(total) else None
        if t is None or similar(r['name'], t['name']) < 0.8 or r['number'] in outlays:
            raise ValueError(f'Treasury Table 5 item {r["number"]} {r["name"]!r} does not match Table 1')
        outlays[r['number']] = r['values']
    flagged = {t['number'] for t in total if '1/' in t['footnotes']}
    if flagged ^ set(outlays):
        differ = [str(n) for n in sorted(flagged ^ set(outlays))]
        listed = differ[0] if len(differ) == 1 else ', '.join(differ[:-1]) + ' and ' + differ[-1]
        issues.append(f'Table 1 footnote 1/ ("see Table 5") and Table 5 coverage differ for provision{"s" if len(differ) > 1 else ""} {listed}; outlays are taken from Table 5 as published.')
    reliefs = []
    for t in total:
        corp_any, ind_any = any(t['corporate']), any(t['individual'])
        head = HEADS[('C' if corp_any else '') + ('I' if ind_any else '')]
        outlay = outlays.get(t['number'])
        estimates = []
        for i, year in enumerate(years):
            forecast = year >= first_forecast
            estimates.append({'year': year, 'fiscalYear': f'FY{int(year[:4]) + 1}', 'costMillion': t['values'][i], 'status': 'forecast' if forecast else 'estimate',
                              'rawValue': str(int(t['values'][i])), 'claimants': None, 'claimantStatus': 'unavailable', 'isForecast': forecast,
                              'corporateMillion': t['corporate'][i], 'individualMillion': t['individual'][i],
                              'outlayMillion': outlay[i] if outlay else None})
        own_notes = [f'{m} {notes[m]}' for m in t['footnotes'] if m in notes and m != '1/']
        if outlay:
            own_notes.insert(0, 'The outlay effect of the refundable portion (Treasury Table 5) is shown separately and is not included in the revenue-loss estimate.')
        if t['number'] in derived:
            own_notes.append('The split between corporate and individual income tax is derived: this provision is missing from one of Treasury Tables 2a and 2b.')
        title, body = descriptions[t['number']]
        reliefs.append({'id': str(t['number']), 'provisionNumber': t['number'], 'name': t['name'], 'taxHead': head,
                        'classification': 'tax expenditure', 'estimateType': 'multi-year', 'objective': None,
                        'description': body or None, 'descriptionTitle': title, 'notes': ' '.join(own_notes) or None,
                        'quality': 'Normal tax method: a tax expenditure only under the normal tax baseline (zero under the reference tax law baseline).' if 'normal tax method' in t['name'].lower() else None,
                        'functionalCategory': t['function'], 'budgetFunction': t['function'], 'tableSubheading': t['subheading'],
                        'normalTaxMethod': 'normal tax method' in t['name'].lower(), 'hasOutlayEffect': outlay is not None,
                        'splitDerived': t['number'] in derived, 'taxExpenditureType': None,
                        'sourceUrl': xlsx_url, 'publicationVintage': published, 'retrievalDate': retrieved[:10] if retrieved else None, 'firstForecastYear': first_forecast,
                        'aggregationRole': 'standalone', 'estimates': estimates})

    def totals(items):
        out = []
        for i, year in enumerate(years):
            values = [r['estimates'][i]['costMillion'] for r in items]
            out.append({'year': year, 'quantifiedSumMillion': round(sum(values), 6), 'quantifiedCount': len(values), 'excluded': {}, 'isForecast': year >= first_forecast})
        return out

    retirement = [r for r in reliefs if r['tableSubheading'] == RETIREMENT]
    names = {r['name'] for r in retirement}
    if not 4 <= len(retirement) <= 12 or not {'Defined benefit employer plans', 'Defined contribution employer plans'} <= names or not any('Individual Retirement Accounts' in n for n in names):
        raise ValueError(f'Treasury retirement grouping ("{RETIREMENT}") changed: {sorted(names)}')
    title = f'Tax Expenditures, Fiscal Year {edition}'
    outlay_totals = [{'year': y, 'outlaySumMillion': round(sum(v[i] for v in outlays.values()), 6), 'count': len(outlays), 'isForecast': y >= first_forecast} for i, y in enumerate(years)]
    data = {'schemaVersion': 1, 'asOf': published, 'publicationDate': published, 'country': 'us', 'edition': f'FY{edition}', 'unit': 'USD million',
            'defaultYear': first_forecast, 'years': years, 'firstForecastYear': first_forecast,
            'sources': [{'id': 'us-treasury-te-workbook', 'name': f'{title} (Excel tables)', 'publisher': PUBLISHER, 'url': xlsx_url, 'publicationDate': published, 'licence': LICENCE, 'attribution': ATTRIBUTION},
                        {'id': 'us-treasury-te-pdf', 'name': f'{title} (PDF: descriptions and tables)', 'publisher': PUBLISHER, 'url': pdf_url, 'publicationDate': published, 'licence': LICENCE, 'attribution': ATTRIBUTION},
                        {'id': 'us-treasury-te-page', 'name': 'Treasury: Tax Expenditures', 'publisher': PUBLISHER, 'url': LANDING, 'publicationDate': published, 'licence': LICENCE, 'attribution': ATTRIBUTION}],
            'aggregate': {'year': first_forecast, 'costMillion': None, 'count': None,
                          'label': 'No official total: Treasury does not publish a grand total and cautions against adding tax expenditures, because provisions interact.',
                          'source': pdf_url, 'rounding': 'A derived, clearly marked sum of the published estimates is shown separately.'},
            'derivedTotals': {'basis': 'derived', 'method': 'Sum of the published receipt (revenue-loss) effects in Treasury Table 1 for each fiscal year, positive and negative; outlay effects (Table 5) are excluded.',
                              'caveat': 'Not an official total and not the revenue from repealing everything: provisions interact (for example itemized deductions jointly versus singly, or exclusions pushing income into higher brackets).',
                              'all': totals(reliefs), 'byTaxHead': {h: totals([r for r in reliefs if r['taxHead'] == h]) for h in dict.fromkeys(r['taxHead'] for r in reliefs)},
                              'outlays': outlay_totals},
            'retirement': {'basis': f'Treasury Table 1 group "{RETIREMENT}" (provisions {retirement[0]["id"]}-{retirement[-1]["id"]}: employer defined benefit and defined contribution plans, IRAs, self-employed plans and related credits)',
                           'codes': [r['id'] for r in retirement], 'totals': totals(retirement),
                           'note': 'Derived sum of the receipt effects of the retirement-saving provisions; they interact, so the sum is descriptive, not a repeal estimate. Social Security benefit exclusions are not included.'},
            'superannuation': None, 'reliefs': reliefs, 'isa': None, 'pensions': None,
            'sourceIssues': issues,
            'notes': ['Tax expenditure estimates are receipt (revenue-loss) effects of each provision against a comprehensive-income baseline, holding the rest of the tax code and taxpayer behavior unchanged. They are not the revenue gain from repeal and must not be added to outlays or the deficit.',
                      'Treasury publishes no grand total: provisions interact, so the sum of individual estimates differs from the effect of repealing them together. Any sum shown here is a derived aggregate, marked as such.',
                      f'Fiscal years run October to September and are named by the year in which they end. FY{int(years[0][:4]) + 1} is an estimate for the latest completed fiscal year at publication ({long_date(published)}); years from FY{int(first_forecast[:4]) + 1} are projections. All figures reflect law as enacted when the estimates were prepared and the Administration economic assumptions then current.',
                      'US$ million, published in multiples of $10 million. Zero is shown as published: Treasury does not distinguish true zeros from amounts that round to zero. Negative values are timing effects (deferrals) and are kept.',
                      'Only income tax provisions are covered. Each provision is classed as individual, corporate, or both according to which of Treasury Tables 2b (individual) and 2a (corporate) show nonzero amounts, and the split between the two is kept. Provisions with no receipt effect in any year are classed as both.',
                      'Refundable credits (for example the earned income tax credit, child tax credit and premium tax credit) also raise outlays. Those outlay effects (Treasury Table 5) are shown separately and are not included in the revenue-loss estimates or in any derived total.',
                      'Each provision is grouped under the federal budget function in which Treasury lists it. Provisions marked "normal tax method" are tax expenditures only under the normal tax baseline.',
                      'Provision numbers change between editions. Descriptions are taken from the PDF text; words split by the text extraction are rejoined where the document confirms the word, but a few may still show a stray space.']}
    validate_reliefs(data)
    return data


def validate_reliefs(data):
    if data.get('schemaVersion') != 1 or data.get('country') != 'us' or data.get('unit') != 'USD million':
        raise ValueError('Unsupported US reliefs schema')
    rows = data.get('reliefs') or []
    years = data.get('years') or []
    if len(rows) < 150:
        raise ValueError('Incomplete Treasury tax expenditure dataset')
    if len(years) != 11 or any(int(b[:4]) - int(a[:4]) != 1 or not re.fullmatch(r'\d{4}-\d{2}', a) for a, b in zip(years, years[1:])):
        raise ValueError('Treasury years are not 11 contiguous fiscal years')
    if data.get('defaultYear') not in years or data['aggregate'].get('year') != data['defaultYear']:
        raise ValueError('Default year outside the published window')
    if len({r['id'] for r in rows}) != len(rows) or len({(r['name'], r['taxHead']) for r in rows}) != len(rows):
        raise ValueError('Duplicate tax expenditure identity')
    for r in rows:
        if r['classification'] != 'tax expenditure' or r.get('taxHead') not in set(HEADS.values()) or not r.get('functionalCategory') or [e['year'] for e in r['estimates']] != years:
            raise ValueError(f'Malformed tax expenditure {r.get("id")}')
        for e in r['estimates']:
            if e['status'] not in STATUSES or e['isForecast'] != (e['year'] >= data['defaultYear']) or e['status'] != ('forecast' if e['isForecast'] else 'estimate'):
                raise ValueError(f'Bad estimate status for {r["id"]} {e["year"]}')
            values = [e['costMillion'], e['corporateMillion'], e['individualMillion']]
            if any(not isinstance(v, float) for v in values) or abs(values[1] + values[2] - values[0]) > 1e-6:
                raise ValueError(f'Non-numeric or inconsistent estimate for {r["id"]} {e["year"]}')
            if abs(e['costMillion']) > 1_000_000 or (e['outlayMillion'] is not None and not 0 <= e['outlayMillion'] < 500_000):
                raise ValueError(f'Implausible magnitude for {r["id"]} {e["year"]}')
    i = years.index(data['defaultYear'])
    top = sorted((r['estimates'][i]['costMillion'] for r in rows), reverse=True)
    if not 100_000 <= top[0] <= 800_000 or top[2] < 50_000:
        raise ValueError(f'Implausible largest tax expenditures: {top[:3]}')
    if data.get('aggregate', {}).get('costMillion') is not None:
        raise ValueError('Treasury publishes no tax expenditure total')
    for i, total in enumerate(data['derivedTotals']['all']):
        values = [r['estimates'][i]['costMillion'] for r in rows]
        if total['year'] != years[i] or abs(sum(values) - total['quantifiedSumMillion']) > 1e-6 or total['quantifiedCount'] != len(rows):
            raise ValueError('Derived Treasury total inconsistent')
        if not 500_000 < total['quantifiedSumMillion'] < 10_000_000:
            raise ValueError('Implausible derived Treasury total')
    codes = data['retirement']['codes']
    if not codes or not set(codes) <= {r['id'] for r in rows}:
        raise ValueError('Retirement codes do not match provisions')
    if data.get('isa') or data.get('pensions') or data.get('superannuation'):
        raise ValueError('UK/Australia-only relief sections must be null')
    return data


def fetch_reliefs(fetch):
    edition, xlsx_url, pdf_url = links(fetch(LANDING).decode('utf-8', 'replace'))
    xlsx, pdf = fetch(xlsx_url), fetch(pdf_url)
    record = next((r for r in reversed(getattr(fetch, 'records', [])) if r['url'] == xlsx_url), {})
    data = parse_reliefs(xlsx, pdf_text(pdf), edition, xlsx_url, pdf_url, record.get('retrievedAt'))
    data['fileSha256'] = {'workbook': hashlib.sha256(xlsx).hexdigest(), 'pdf': hashlib.sha256(pdf).hexdigest()}
    return data
