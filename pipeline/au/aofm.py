"""AOFM Australian Government Securities: stock on issue, tender results and next-week calendar.

Stock: weekly HTML on-issue tables (face value, AUD million, ISIN) for Treasury Bonds,
Treasury Indexed Bonds and Treasury Notes. TIB face value is unindexed; upliftedMillion
for TIBs is DERIVED as face x Kt/100 using the AOFM capital indexation factor for the
next coupon date. Tender and syndication results come from the Data Hub issuance
workbooks (full AUD, converted to millions). The calendar is the AOFM's one-week-ahead
Forthcoming Transactions page; past weeks are overwritten at source, so only the current
snapshot is available. Data Hub file paths embed upload dates and move, so each run
discovers them from the Data Hub page. aofm.gov.au tarpits requests that lack browser-like
Accept/Accept-Language/Accept-Encoding headers; every request sends them.
"""
from datetime import date, datetime, timedelta, timezone
from html.parser import HTMLParser
from io import BytesIO
from itertools import accumulate
import json
import math
from pathlib import Path
import re
from urllib.parse import unquote, urljoin

CATALOGUE = Path(__file__).resolve().parents[2] / 'catalogue' / 'au' / 'aofm.json'
BASE = 'https://www.aofm.gov.au'
HEADERS = {'User-Agent': 'uk-fiscal-research/0.1 (research tool)',
           'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
           'Accept-Language': 'en-AU,en;q=0.9', 'Accept-Encoding': 'gzip, deflate'}
TB, TIB, TN = 'Treasury Bond', 'Treasury Indexed Bond', 'Treasury Note'
DATE = r'\d{1,2} (?:January|February|March|April|May|June|July|August|September|October|November|December) \d{4}'
HUB_FILES = {  # Data Hub link text (case-insensitive) -> key; an _N suffix marks a re-upload.
    'tb_issuance': r'treasury bonds - issuance(?:_\d+)?\.xlsx',
    'tib_issuance': r'treasury indexed bonds - issuance(?:_\d+)?\.xlsx',
    'tn_issuance': r'treasury notes - issuance(?:_\d+)?\.xlsx',
    'portfolio': r'portfolio_aggregate_-_dealt(?:_\d+)?\.xlsx',
    'indexation': r'treasury_indexed_bonds_-_indexation_factors(?:_\d+)?\.xlsx',
    'foreign': r'foreign_holdings(?:_\d+)?\.xlsx',
}
AUCTION_YEARS = 5


def clean(s):
    return ' '.join(str(s).replace('\xa0', ' ').split())


def number(s):
    n = float(str(s).replace(',', '').replace('$', '').strip())
    if not math.isfinite(n):
        raise ValueError('Non-finite AOFM value')
    return n


def day(s):
    s = clean(s)
    m = re.search(DATE, s)
    if not m:
        raise ValueError('Unrecognised AOFM date: ' + s)
    return datetime.strptime(m[0], '%d %B %Y').date()


def pct_text(c):
    s = f'{c:.4f}'.rstrip('0')
    return s + '0' * max(0, 2 - len(s.split('.')[1]))


def label(instrument, maturity, coupon=None, green=False):
    when = f'{maturity.day} {maturity:%B %Y}'
    if instrument == TN:
        return f'{TN} {when}'
    return f"{pct_text(coupon)}% {when} {'Green ' if green else ''}{instrument}"


class Page(HTMLParser):
    """Tables (with colspan expanded), the h2 they sit under, section text and links."""
    def __init__(self, raw):
        super().__init__(convert_charrefs=True)
        self.tables, self.links, self.sections, self.items = [], [], {None: ''}, []
        self.heading, self.h2, self.cell, self.link, self.li = None, None, None, None, None
        self.feed(raw.decode('utf-8', 'replace') if isinstance(raw, bytes) else raw)
        self.text = clean(' '.join(self.sections.values()))

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == 'h2':
            self.h2 = ''
        elif tag == 'table':
            self.tables.append({'heading': self.heading, 'rows': []})
        elif tag == 'tr' and self.tables:
            self.tables[-1]['rows'].append([])
        elif tag in ('td', 'th'):
            self.cell = ['', int(a.get('colspan') or 1)]
        elif tag == 'a':
            self.link = [a.get('href') or '', '']
        elif tag == 'li':
            self.li = ''

    def handle_endtag(self, tag):
        if tag == 'h2' and self.h2 is not None:
            self.heading = clean(self.h2)
            self.sections.setdefault(self.heading, '')
            self.h2 = None
        elif tag in ('td', 'th') and self.cell and self.tables and self.tables[-1]['rows']:
            self.tables[-1]['rows'][-1].extend([clean(self.cell[0])] * self.cell[1])
            self.cell = None
        elif tag == 'a' and self.link:
            self.links.append((self.link[0], clean(self.link[1])))
            self.link = None
        elif tag == 'li' and self.li is not None:
            self.items.append(clean(self.li))
            self.li = None

    def handle_data(self, data):
        for buf in ('h2', 'link', 'li'):
            if getattr(self, buf) is not None:
                if buf == 'link':
                    self.link[1] += data
                else:
                    setattr(self, buf, getattr(self, buf) + data)
        if self.cell is not None:
            self.cell[0] += data
        elif self.h2 is None:
            self.sections[self.heading] += ' ' + data


def discover(raw):
    """Current Data Hub file URLs, matched by link text; exactly one match per file."""
    found = {}
    for href, text in Page(raw).links:
        name = clean(text or unquote(href.rsplit('/', 1)[-1])).lower()
        for key, pattern in HUB_FILES.items():
            if re.fullmatch(pattern, name) and '/sites/default/files/' in href:
                found.setdefault(key, set()).add(urljoin(BASE + '/', href))
    for key in HUB_FILES:
        if len(found.get(key, ())) != 1:
            raise ValueError(f'AOFM Data Hub link for {key} not found exactly once: {sorted(found.get(key, ()))}')
    return {k: v.pop() for k, v in found.items()}


def parse_on_issue(raw, instrument, green_isins=()):
    page = Page(raw)
    m = re.search(r'on issue as(?: at)? (' + DATE + ')', page.text)
    if not m:
        raise ValueError(f'{instrument} on-issue date not found')
    as_of = day(m[1])
    tables = [t for t in page.tables if t['rows'] and any('ISIN' in c for c in t['rows'][0])]
    if len(tables) != 1:
        raise ValueError(f'Expected one {instrument} on-issue table')
    head = tables[0]['rows'][0]
    if len(head) != 3 or 'face value' not in head[1].lower() or 'million' not in head[1].lower():
        raise ValueError(f'{instrument} on-issue table layout changed: {head}')
    rows = []
    for name, amount, isin in (r for r in tables[0]['rows'][1:] if r):
        isin = isin.strip()
        if not re.fullmatch(r'AU[0-9A-Z]{10}', isin):
            raise ValueError('Malformed AGS ISIN: ' + isin)
        if instrument == TN:
            coupon = 0.0
        else:
            c = re.fullmatch(r'([\d.]+)% (' + DATE + ')', name)
            if not c:
                raise ValueError(f'Unrecognised {instrument} line: {name}')
            coupon = number(c[1])
        maturity = day(name)
        green = isin in green_isins
        rows.append({'isin': isin, 'name': label(instrument, maturity, coupon, green),
                     'type': 'index-linked' if instrument == TIB else 'conventional',
                     'instrument': instrument, 'green': green, 'couponPct': coupon,
                     'maturityDate': maturity.isoformat(), 'nominalMillion': number(amount)})
    if not rows:
        raise ValueError(f'No {instrument} lines on issue')
    return as_of.isoformat(), rows


def sheet_rows(raw, name):
    import openpyxl
    wb = openpyxl.load_workbook(BytesIO(raw), read_only=True, data_only=True)
    if name not in wb.sheetnames:
        raise ValueError(f'AOFM workbook lacks sheet {name}: {wb.sheetnames}')
    return [list(r) for r in wb[name].iter_rows(values_only=True)]


def blank(v):
    return v is None or (isinstance(v, str) and not v.strip())


def parse_kt(raw):
    """(maturity, coupon) -> sorted [(coupon date, Kt)] from the KtFactors sheet."""
    rows = sheet_rows(raw, 'KtFactors')
    find = lambda first: next((r for r in rows if r and clean(r[0] or '') == first), None)
    coupons, maturities, head = find('Coupon (%)'), find('Maturity'), find('CPI Date')
    if not (coupons and maturities and head):
        raise ValueError('Indexation factor layout changed')
    start = rows.index(head) + 1
    series = {}
    for col in range(1, len(head), 3):
        if head[col:col + 2] != ['Coupon Date', 'Kt'] or not isinstance(maturities[col], datetime):
            continue
        key = (maturities[col].date().isoformat(), round(float(coupons[col]), 4))
        series[key] = sorted((r[col].date(), number(r[col + 1])) for r in rows[start:]
                             if isinstance(r[col], datetime) and not blank(r[col + 1]))
    if not series:
        raise ValueError('No indexation factors')
    return series


def apply_uplift(securities, kt, as_of):
    ref = date.fromisoformat(as_of)
    for s in securities:
        if s['instrument'] != TIB:
            s['upliftedMillion'] = s['nominalMillion']
            continue
        factors = kt.get((s['maturityDate'], round(s['couponPct'], 4)))
        nxt = [(d, k) for d, k in factors or () if d > ref]
        if not nxt:
            raise ValueError('No current indexation factor for ' + s['name'])
        d, k = nxt[0]
        s['upliftedMillion'] = round(s['nominalMillion'] * k / 100, 6)
        s['indexation'] = {'kt': k, 'couponDate': d.isoformat(), 'derived': True}


def parse_issuance(raw, instrument, source_url, green_isins=()):
    rows = sheet_rows(raw, 'Transactions')
    hi = next(i for i, r in enumerate(rows) if r and r[0] == 'Date Held')
    col = {clean(h): i for i, h in enumerate(rows[hi]) if not blank(h)}
    need = ['Date Held', 'Tender Number', 'Maturity', 'ISIN', 'Amount Offered', 'Amount Allotted', 'Amount of Bids',
            'Coverage Ratio', 'Weighted Average Issue Yield', 'Lowest Accepted Yield', 'Highest Accepted Yield', 'Date Settled']
    need += [] if instrument == TN else ['Coupon', 'Secondary Market Mid Rate']
    missing = [n for n in need if n not in col]
    if missing:
        raise ValueError(f'{instrument} issuance workbook lacks columns {missing}')
    units = rows[hi + 1]
    if any(clean(units[col[n]] or '') != '$' for n in ('Amount Offered', 'Amount Allotted', 'Amount of Bids')):
        raise ValueError(f'{instrument} issuance amounts are no longer in full dollars')
    get = lambda r, n: None if n not in col or blank(r[col[n]]) else r[col[n]]
    mill = lambda v: None if v is None else round(number(v) / 1e6, 6)
    num = lambda v: None if v is None else number(v)
    out = []
    for r in rows[hi + 2:]:
        if not r or not isinstance(r[0], datetime):
            if r and any(not blank(v) for v in r):
                raise ValueError(f'Unexpected {instrument} issuance row: {r[:5]}')
            continue
        tender = clean(get(r, 'Tender Number') or '')
        prefix = tender.split('/')[-1].rstrip('0123456789').strip() if instrument != TN else ''
        method = 'syndication' if 'SYN' in prefix else 'tender' if prefix in ('', 'TB', 'TIB', 'INDEXTENDER') else 'other'
        isin = clean(get(r, 'ISIN') or '') or None
        green = tender.startswith('Green Bond/') or (isin in green_isins)
        maturity = get(r, 'Maturity')
        if not isinstance(maturity, datetime):
            raise ValueError(f'{instrument} issuance row lacks maturity: {r[:5]}')
        coupon = None if instrument == TN else num(get(r, 'Coupon'))
        wavg, high = num(get(r, 'Weighted Average Issue Yield')), num(get(r, 'Highest Accepted Yield'))
        settled = get(r, 'Date Settled')
        out.append({'date': r[0].date().isoformat(), 'isin': isin,
                    'name': label(instrument, maturity.date(), coupon, green),
                    'type': 'index-linked' if instrument == TIB else 'conventional',
                    'instrument': instrument, 'method': method, 'tenderNumber': tender, 'green': green,
                    'maturityDate': maturity.date().isoformat(), 'couponPct': coupon,
                    'offeredMillion': mill(get(r, 'Amount Offered')), 'allottedMillion': mill(get(r, 'Amount Allotted')),
                    'bidsMillion': mill(get(r, 'Amount of Bids')), 'yieldPct': wavg,
                    'yieldBasis': ('real' if instrument == TIB else 'nominal') + (' syndication yield' if method == 'syndication' else ' weighted average issue yield'),
                    'cover': num(get(r, 'Coverage Ratio')),
                    'tailBp': round((high - wavg) * 100, 4) if method == 'tender' and high is not None and wavg is not None else None,
                    'lowestAcceptedYieldPct': num(get(r, 'Lowest Accepted Yield')), 'highestAcceptedYieldPct': high,
                    'secondaryMidPct': num(get(r, 'Secondary Market Mid Rate')),
                    'postAuctionMillion': None,
                    'settlementDate': settled.date().isoformat() if isinstance(settled, datetime) else None,
                    'sourceUrl': source_url, 'publicationDate': None})
    if not out:
        raise ValueError(f'No {instrument} issuance rows')
    return out


HISTORY_COLUMNS = ['date', 'instrument', 'method', 'tenderNumber', 'isin', 'maturityDate', 'couponPct',
                   'offeredMillion', 'allottedMillion', 'bidsMillion', 'cover', 'yieldPct']


def compact(rows):
    short = {TB: 'TB', TIB: 'TIB', TN: 'TN'}
    return {'columns': HISTORY_COLUMNS, 'instrumentCodes': {v: k for k, v in short.items()},
            'rows': [[short[r[c]] if c == 'instrument' else r[c] for c in HISTORY_COLUMNS] for r in rows]}


def parse_forthcoming(raw, url, green_isins=()):
    """Tender tables are transposed: a th label per row, one column per line."""
    page = Page(raw)
    kinds = {'Treasury Bonds': TB, 'Treasury Indexed Bonds': TIB, 'Treasury Notes': TN}
    events = []
    for t in page.tables:
        instrument = kinds.get(t['heading'])
        if instrument is None:
            raise ValueError('Forthcoming transactions table under unrecognised heading: ' + str(t['heading']))
        fields = {r[0]: r[1:] for r in t['rows'] if r}
        series = fields.get('Maturity Date' if instrument == TN else 'Series Offered')
        amount = next((v for k, v in fields.items() if k.startswith('Offered to Public')), None)
        if not series or amount is None or not all(k in fields for k in ('ISIN', 'Tender Date', 'Settlement Date')):
            raise ValueError(f'Forthcoming {instrument} table layout changed: {list(fields)}')
        if '$million' not in next(k for k in fields if k.startswith('Offered to Public')).replace(' ', ''):
            raise ValueError('Forthcoming amounts no longer in $million')
        for i, name in enumerate(series):
            isin = fields['ISIN'][i].strip()
            maturity = day(name)
            coupon = None if instrument == TN else number(re.match(r'([\d.]+)%', name)[1])
            events.append({'date': day(fields['Tender Date'][i]).isoformat(), 'type': 'tender',
                           'name': label(instrument, maturity, coupon, isin in green_isins),
                           'instrument': instrument, 'isin': isin, 'amountMillion': number(amount[i]),
                           'settlementDate': day(fields['Settlement Date'][i]).isoformat(),
                           'bidWindow': (fields.get('Time to Submit Bids') or [None] * len(series))[i],
                           'announcementDate': None, 'status': 'announced', 'datePrecision': 'day',
                           'sourceUrl': url, 'publicationDate': None})
    for heading, instrument in kinds.items():
        if any(t['heading'] == heading for t in page.tables):
            continue
        m = re.search(r'next tender .*?planned to be held on (?:\w+, )?(' + DATE + ')', page.sections.get(heading, ''))
        if m:
            events.append({'date': day(m[1]).isoformat(), 'type': 'tender',
                           'name': f'{instrument} tender (line and size not yet announced)', 'instrument': instrument,
                           'isin': None, 'amountMillion': None, 'settlementDate': None, 'bidWindow': None,
                           'announcementDate': None, 'status': 'planned', 'datePrecision': 'day',
                           'sourceUrl': url, 'publicationDate': None})
    if not any(h in page.sections for h in kinds):
        raise ValueError('Forthcoming transactions page has no instrument sections')
    return sorted(events, key=lambda e: (e['date'], e['name']))


def parse_program(raw, url):
    """Latest qualitative program update: date, title and guidance bullets."""
    page = Page(raw)
    m = re.search(DATE, page.text)
    items = [i for i in page.items if re.match(r'(?i)(issue|conduct|hold|maintain)\b', i)]
    if not m or not items:
        raise ValueError('Issuance program page layout changed')
    year = re.search(r'(\d{4}-\d{2}) issuance program', page.text)
    return {'publicationDate': day(m[0]).isoformat(), 'fiscalYear': year[1] if year else None,
            'guidance': items, 'sourceUrl': url}


def latest_headline_link(raw):
    for href, text in Page(raw).links:
        if re.search(r'(?i)issuance program.*(budget|myefo)', text):
            return urljoin(BASE + '/', href), text
    raise ValueError('No Budget/MYEFO issuance program update linked')


def parse_headline(raw, url, title):
    """Headline planned issuance in a Budget/MYEFO notice. Unmatched amounts stay null."""
    text = Page(raw).text
    m = re.search(DATE, text)
    tb = re.search(r'Treasury Bonds(?: \(including Green Treasury Bonds\))? (?:in|for) (\d{4}-\d{2}) '
                   r'(?:is expected to be|has been revised to|remains) around \$([\d.]+) billion', text)
    tib = re.search(r'Treasury Indexed Bonds(?: by tender)?(?: in \d{4}-\d{2})? (?:is expected to be|remains) '
                    r'(?:between \$([\d.]+) billion and \$([\d.]+) billion|around \$([\d.]+) billion)', text)
    return {'title': title, 'publicationDate': day(m[0]).isoformat() if m else None,
            'fiscalYear': tb[1] if tb else None,
            'treasuryBondsBillion': number(tb[2]) if tb else None,
            'treasuryIndexedBondsLowBillion': number(tib[1] or tib[3]) if tib else None,
            'treasuryIndexedBondsHighBillion': number(tib[2] or tib[3]) if tib else None,
            'basis': 'Planned gross issuance stated in the notice text ("around"/"between"); null where the wording was not recognised.',
            'sourceUrl': url}


def parse_portfolio(raw):
    """Month-end face value of AUD TB, TIB (unindexed) and TN liabilities; signs flipped to positive."""
    rows = sheet_rows(raw, 'FaceValue')
    labels = {clean(r[0] or ''): r for r in rows[:8] if r}
    side, cat, inst = labels.get('Liability / Asset'), labels.get('Category'), labels.get('Instrument')
    cur = next((r for k, r in labels.items() if k.startswith('Currency of Issue')), None)
    if not (side and cat and inst and cur):
        raise ValueError('Portfolio aggregate header layout changed')
    fill = lambda r: list(accumulate(r, lambda a, b: b if not blank(b) else a))
    side, cat, cur = fill(side), fill(cat), fill(cur)
    want = {'Treasury Bonds': 'treasuryBondsMillion', 'Treasury Indexed Bonds (Capital Indexed)': 'treasuryIndexedBondsMillion',
            'Treasury Notes': 'treasuryNotesMillion'}
    cols = {want[clean(inst[i])]: i for i in range(1, len(inst))
            if clean(inst[i] or '') in want and side[i] == 'LIABILITY' and cur[i] == 'AUD' and cat[i] == 'Main Funding Instruments'}
    if len(cols) != 3:
        raise ValueError('Portfolio aggregate instrument columns not found')
    out = []
    for r in rows:
        if r and isinstance(r[0], datetime):
            out.append({'date': r[0].date().isoformat(),
                        **{k: None if blank(r[i]) else round(-number(r[i]) / 1e6, 6) for k, i in cols.items()}})
    if not out:
        raise ValueError('No portfolio aggregate observations')
    return out


def parse_foreign(raw):
    rows = sheet_rows(raw, 'AGS')
    hi = next(i for i, r in enumerate(rows) if r and 'AGS Held by Non-Resident Holdings' in [clean(v or '') for v in r])
    col = {clean(v or ''): i for i, v in enumerate(rows[hi])}
    need = ['AGS Held by Non-Resident Holdings', 'Total AGS Outstanding', 'Proportion Held by Non-Residents',
            'Proportion held by Non-Residents Adjusted for Net Repo Outflow']
    if any(n not in col for n in need):
        raise ValueError('Foreign holdings layout changed')
    out = []
    for r in rows[hi + 1:]:
        if r and isinstance(r[0], datetime) and not blank(r[col[need[2]]]):
            share = number(r[col[need[2]]])
            if not 0 <= share <= 1:
                raise ValueError('Non-resident share outside 0-1')
            out.append({'date': r[0].date().isoformat(),
                        'heldMillion': round(number(r[col[need[0]]]) / 1e6, 3),
                        'outstandingMillion': round(number(r[col[need[1]]]) / 1e6, 3),
                        'sharePct': round(share * 100, 4),
                        'repoAdjustedSharePct': None if blank(r[col[need[3]]]) else round(number(r[col[need[3]]]) * 100, 4)})
    if not out:
        raise ValueError('No foreign holdings observations')
    return out


def reconcile(securities, as_of, history, auctions):
    """Weekly on-issue total vs latest month-end aggregate plus allotments dealt since. Informational only."""
    last = history[-1]
    out = []
    for inst, key in ((TB, 'treasuryBondsMillion'), (TIB, 'treasuryIndexedBondsMillion'), (TN, 'treasuryNotesMillion')):
        web = sum(s['nominalMillion'] for s in securities if s['instrument'] == inst)
        dealt = sum(a['allottedMillion'] or 0 for a in auctions if a['instrument'] == inst and last['date'] < a['date'] <= as_of)
        base = last[key]
        out.append({'instrument': inst, 'monthEnd': last['date'], 'monthEndMillion': base,
                    'allottedSinceMillion': round(dealt, 6), 'onIssueMillion': round(web, 6),
                    'residualMillion': None if base is None else round(web - base - dealt, 6)})
    return out


def fetch_debt(fetch):
    cat = json.loads(CATALOGUE.read_text())
    now = datetime.now(timezone.utc).isoformat()
    get = lambda url: fetch(url, headers=HEADERS)
    pages = cat['pages']
    hub = discover(get(pages['dataHub']))
    green = set(cat['greenIsins'])
    tb_rows = parse_issuance(get(hub['tb_issuance']), TB, hub['tb_issuance'], green)
    green |= {a['isin'] for a in tb_rows if a['green'] and a['isin']}
    dates, securities = {}, []
    for inst, page in ((TB, 'treasuryBonds'), (TIB, 'treasuryIndexedBonds'), (TN, 'treasuryNotes')):
        d, rows = parse_on_issue(get(pages[page]), inst, green)
        dates[inst] = d
        securities += rows
    as_of = min(dates.values())
    apply_uplift(securities, parse_kt(get(hub['indexation'])), as_of)
    securities.sort(key=lambda s: (s['maturityDate'], s['isin']))
    history = tb_rows + parse_issuance(get(hub['tib_issuance']), TIB, hub['tib_issuance'], green) \
        + parse_issuance(get(hub['tn_issuance']), TN, hub['tn_issuance'], green)
    history.sort(key=lambda a: (a['date'], a['instrument'], a['maturityDate']))
    cutoff = (date.fromisoformat(as_of) - timedelta(days=round(365.25 * AUCTION_YEARS))).isoformat()
    auctions = [a for a in history if a['date'] > cutoff and a['method'] in ('tender', 'syndication')]
    calendar = parse_forthcoming(get(pages['forthcoming']), pages['forthcoming'], green)
    program = parse_program(get(pages['issuanceProgram']), pages['issuanceProgram'])
    head_url, head_title = latest_headline_link(get(pages['historicalProgram']))
    program['headline'] = parse_headline(get(head_url), head_url, head_title)
    stock_history = parse_portfolio(get(hub['portfolio']))
    foreign = parse_foreign(get(hub['foreign']))
    by_inst = {}
    for s in securities:
        t = by_inst.setdefault(s['instrument'], {'lines': 0, 'nominalMillion': 0.0, 'upliftedMillion': 0.0, 'observationDate': dates[s['instrument']]})
        t['lines'] += 1
        t['nominalMillion'] = round(t['nominalMillion'] + s['nominalMillion'], 6)
        t['upliftedMillion'] = round(t['upliftedMillion'] + s['upliftedMillion'], 6)
    meta = {k: cat[k] for k in ('publisher', 'licence', 'attribution', 'licenceUrl')}
    src = lambda id, name, url, obs, **kw: {'id': id, 'name': name, 'title': name, **meta, 'url': url,
                                             'landingUrl': kw.pop('landing', pages['dataHub']), 'observationDate': obs,
                                             'publicationDate': None, 'retrievedAt': now, **kw}
    sources = [
        src('aofm-tb-on-issue', 'Treasury Bonds on issue', pages['treasuryBonds'], dates[TB], landing=pages['treasuryBonds'], units='AUD million face value', frequency='Weekly'),
        src('aofm-tib-on-issue', 'Treasury Indexed Bonds on issue', pages['treasuryIndexedBonds'], dates[TIB], landing=pages['treasuryIndexedBonds'], units='AUD million face value (unindexed)', frequency='Weekly'),
        src('aofm-tn-on-issue', 'Treasury Notes on issue', pages['treasuryNotes'], dates[TN], landing=pages['treasuryNotes'], units='AUD million face value', frequency='Weekly'),
        src('aofm-tib-indexation', 'Treasury Indexed Bonds - indexation factors (Kt)', hub['indexation'], None, units='Index, base 100 at issue'),
        src('aofm-tb-issuance', 'Treasury Bonds - Issuance', hub['tb_issuance'], max(a['date'] for a in history if a['instrument'] == TB), units='Full AUD (published); converted to AUD million'),
        src('aofm-tib-issuance', 'Treasury Indexed Bonds - Issuance', hub['tib_issuance'], max(a['date'] for a in history if a['instrument'] == TIB), units='Full AUD (published); converted to AUD million'),
        src('aofm-tn-issuance', 'Treasury Notes - Issuance', hub['tn_issuance'], max(a['date'] for a in history if a['instrument'] == TN), units='Full AUD (published); converted to AUD million'),
        src('aofm-forthcoming', 'Forthcoming Transactions', pages['forthcoming'], None, landing=pages['forthcoming'], frequency='Weekly, next week only'),
        src('aofm-issuance-program', 'Issuance Program Update', pages['issuanceProgram'], None, landing=pages['issuanceProgram'], publicationDate=program['publicationDate']),
        src('aofm-issuance-headline', head_title, head_url, None, landing=pages['historicalProgram'], publicationDate=program['headline']['publicationDate']),
        src('aofm-portfolio-aggregate', 'Portfolio Aggregate (dealt basis), FaceValue', hub['portfolio'], stock_history[-1]['date'], units='Full AUD (published); converted to AUD million', frequency='Monthly'),
        src('aofm-foreign-holdings', 'Non-resident holdings of AGS', hub['foreign'], foreign[-1]['date'], units='Full AUD market value; converted to AUD million', frequency='Quarterly'),
    ]
    return {'schemaVersion': 1, 'asOf': as_of, 'retrievedAt': now, 'sources': sources,
            'securities': securities,
            'totals': {'nominalMillion': round(sum(s['nominalMillion'] for s in securities), 6),
                       'upliftedMillion': round(sum(s['upliftedMillion'] for s in securities), 6)},
            'totalsByInstrument': by_inst,
            'auctions': sorted(auctions, key=lambda a: (a['date'], a['instrument'], a['maturityDate']), reverse=True),
            'auctionHistory': compact(history),
            'calendar': calendar,
            'issuanceProgram': program,
            'stockHistory': stock_history,
            'stockReconciliation': reconcile(securities, as_of, stock_history, history),
            'foreignHoldings': foreign,
            'availability': {**cat['availability'], 'auctionsFrom': auctions[-1]['date'] if auctions else None,
                             'auctionHistoryFrom': history[0]['date'], 'lastSourceDiscoveryDate': now[:10]},
            'definitions': cat['definitions'], 'notes': cat['notes']}


def validate_debt(data):
    for key in ('asOf', 'sources', 'securities', 'totals', 'auctions', 'calendar', 'availability', 'definitions', 'notes'):
        if key not in data:
            raise ValueError('Debt dataset lacks ' + key)
    as_of = date.fromisoformat(data['asOf'])
    secs = data['securities']
    if not secs:
        raise ValueError('No debt securities')
    if len({s['isin'] for s in secs}) != len(secs):
        raise ValueError('Duplicate ISIN in securities')
    kinds = {s['instrument'] for s in secs}
    if kinds != {TB, TIB, TN}:
        raise ValueError(f'Securities missing an instrument: {kinds}')
    for s in secs:
        if not re.fullmatch(r'AU[0-9A-Z]{10}', s['isin']):
            raise ValueError('Malformed ISIN ' + s['isin'])
        if date.fromisoformat(s['maturityDate']) <= as_of:
            raise ValueError('Matured line in stock: ' + s['name'])
        if not s['nominalMillion'] > 0 or not s['upliftedMillion'] > 0:
            raise ValueError('Non-positive face value: ' + s['name'])
        if s['type'] == 'conventional' and s['upliftedMillion'] != s['nominalMillion']:
            raise ValueError('Nominal line has an uplift: ' + s['name'])
        if s['type'] == 'index-linked' and not s['upliftedMillion'] >= s['nominalMillion']:
            raise ValueError('Indexed capital value below face: ' + s['name'])
        if (s['type'] == 'index-linked') != (s['instrument'] == TIB):
            raise ValueError('Instrument/type mismatch: ' + s['name'])
    for k in ('nominalMillion', 'upliftedMillion'):
        if abs(sum(s[k] for s in secs) - data['totals'][k]) > .001:
            raise ValueError('Debt total mismatch: ' + k)
    for inst, t in data.get('totalsByInstrument', {}).items():
        if abs(sum(s['nominalMillion'] for s in secs if s['instrument'] == inst) - t['nominalMillion']) > .001:
            raise ValueError('Instrument total mismatch: ' + inst)
    auctions = data['auctions']
    if not auctions:
        raise ValueError('No auction results')
    if len({(a['date'], a['instrument'], a['isin'] or a['maturityDate'], a['tenderNumber']) for a in auctions}) != len(auctions):
        raise ValueError('Duplicate auction result')
    latest = max(a['date'] for a in auctions)
    if date.fromisoformat(latest) < as_of - timedelta(days=45):
        raise ValueError('Auction results are stale relative to the stock snapshot')
    for a in auctions:
        if not (a['allottedMillion'] or 0) > 0:
            raise ValueError(f"Auction without allotment: {a['date']} {a['name']}")
        if a['method'] == 'tender':
            if a['offeredMillion'] is None or a['bidsMillion'] is None or a['cover'] is None or a['yieldPct'] is None:
                raise ValueError(f"Incomplete tender result: {a['date']} {a['name']}")
            if abs(a['bidsMillion'] / a['offeredMillion'] - a['cover']) > .01:
                raise ValueError(f"Cover ratio inconsistent: {a['date']} {a['name']}")
    hist = data.get('auctionHistory')
    if hist and any(len(r) != len(hist['columns']) for r in hist['rows']):
        raise ValueError('Malformed auction history row')
    for e in data['calendar']:
        date.fromisoformat(e['date'])
        if e['amountMillion'] is not None and not e['amountMillion'] > 0:
            raise ValueError('Non-positive announced amount')
    for row in data.get('foreignHoldings', []):
        if not 0 <= row['sharePct'] <= 100:
            raise ValueError('Foreign share out of range')
