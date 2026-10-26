import copy
import io
import unittest
import warnings
from pathlib import Path
import openpyxl
from pipeline.us import composition as c

FIX = Path(__file__).parent / 'fixtures/us/composition'
BASE = 'https://www.whitehouse.gov/wp-content/uploads/2026/04/'
CODES = ['hist01z1', 'hist01z2', 'hist02z1', 'hist02z4', 'hist02z5', 'hist03z1']
FILES = {c.PAGE: 'historical-tables.html', **{f'{BASE}{code}_fy2027.xlsx': f'{code}_fy2027.xlsx' for code in CODES}}


def fetcher(overrides=None):
    calls = []
    def fetch(url, headers=None):
        calls.append(url)
        if not url.startswith('https://'):
            raise AssertionError('HTTPS only')
        if overrides and url in overrides:
            return overrides[url]
        return (FIX / FILES[url]).read_bytes()
    fetch.calls = calls
    return fetch


def edit(code, fn):
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        book = openpyxl.load_workbook(FIX / f'{code}_fy2027.xlsx')
    fn(book.worksheets[0])
    out = io.BytesIO(); book.save(out)
    return {f'{BASE}{code}_fy2027.xlsx': out.getvalue()}


def cell(ws, row_label, col_label, header_row=None):
    """Locate a value cell by row label and year (year-in-row tables when header_row is None)."""
    for row in ws.iter_rows():
        if header_row is None and str(row[0].value) == col_label:
            return row[row_label]
        if header_row is not None and str(row[0].value).strip() == row_label:
            col = next(i for i, v in enumerate(next(ws.iter_rows(min_row=header_row + 1, max_row=header_row + 1, values_only=True))) if str(v) == col_label)
            return row[col]
    raise KeyError(row_label)


def item(row, name):
    return next(i for i in row['items'] if i['name'] == name)


class CompositionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            cls.fetch = fetcher()
            cls.data = c.fetch_composition(cls.fetch)

    def test_edition_and_latest_actual_year(self):
        d = self.data
        self.assertEqual((d['asOf'], d['publicationDate'], d['level'], d['unit']), ('2024-25', '2026-04-03', 'federal', 'USD million'))
        self.assertEqual(d['edition']['budgetYear'], 2027)
        self.assertEqual([r['year'] for r in d['history']['years']], [f'{y}-{str(y + 1)[2:]}' for y in range(2017, 2025)])
        self.assertTrue(all(u.startswith('https://') for u in self.fetch.calls))
        c.validate_composition(d)

    def test_outlays_by_function(self):
        last = self.data['history']['years'][-1]
        self.assertEqual(last['total'], 7011105)
        self.assertEqual(item(last, 'Social Security')['value'], 1580673)
        self.assertEqual(item(last, 'Net interest')['value'], 970065)
        self.assertEqual(item(last, 'Undistributed offsetting receipts')['value'], -150254)
        self.assertEqual(item(last, 'Commerce and housing credit')['value'], -28131)
        self.assertEqual(item(last, 'Other functions')['value'], 20765 + 89834 + 82374 + 45171 + 41966 + 47447 + 83146 + 40546)
        self.assertEqual(sum(i['value'] for i in last['items']), last['total'])
        self.assertEqual(item(self.data['history']['years'][5], 'Education, training, employment and social services')['value'], -2189)

    def test_gdp_shares_use_omb_fiscal_year_gdp(self):
        gdp = {r['year']: r['nominalMillion'] for r in self.data['gdp']['years']}
        self.assertEqual(gdp['2024-25'], 30322847)
        last = self.data['history']['years'][-1]
        self.assertAlmostEqual(last['totalPctGdp'], 7011105 / 30322847 * 100, places=3)
        self.assertEqual(round(last['totalPctGdp'], 1), 23.1)  # OMB Table 1.2.
        self.assertEqual(round(self.data['receipts']['years'][-1]['totalPctGdp'], 1), 17.3)

    def test_receipts_by_source(self):
        r = self.data['receipts']['years'][-1]
        self.assertEqual(r['total'], 5236421)
        self.assertEqual(item(r, 'Social Security payroll taxes (OASDI)')['value'], 1097382 + 186354)
        self.assertEqual(item(r, 'Medicare payroll taxes (hospital insurance)')['value'], 395350)
        self.assertEqual(item(r, 'Customs duties and fees')['value'], 194866)
        self.assertEqual(r['taxTotal'], 5236421 - 5490 - 44239)
        self.assertEqual(sum(i['value'] for i in r['items']), r['total'])

    def test_small_rounding_is_explicit_and_large_breaks_fail(self):
        def bump(by):
            return edit('hist03z1', lambda ws: setattr(cell(ws, 'Total, Federal outlays', '2025', 1), 'value', cell(ws, 'Total, Federal outlays', '2025', 1).value + by))
        def bump11(by):
            return edit('hist01z1', lambda ws: setattr(cell(ws, 2, '2025'), 'value', cell(ws, 2, '2025').value + by))
        data = c.fetch_composition(fetcher({**bump(3), **bump11(3)}))
        self.assertEqual(data['history']['years'][-1]['items'][-1], {'name': c.ROUNDING, 'value': 3, 'pctGdp': round(3 / 30322847 * 100, 4)})
        c.validate_composition(data)
        with self.assertRaisesRegex(ValueError, 'reconciliation'):
            c.fetch_composition(fetcher({**bump(50), **bump11(50)}))
        with self.assertRaisesRegex(ValueError, 'Table 1.1'):
            c.fetch_composition(fetcher(bump11(50)))

    def test_receipt_tables_must_agree(self):
        override = edit('hist02z5', lambda ws: setattr(cell(ws, 2, '2024'), 'value', cell(ws, 2, '2024').value + 100))
        with self.assertRaisesRegex(ValueError, 'Table 2.5'):
            c.fetch_composition(fetcher(override))

    def test_estimates_are_never_treated_as_actuals(self):
        def relabel(ws):
            for row in ws.iter_rows(min_row=2, max_row=2):
                for x in row:
                    if x.value == '2026 estimate':
                        x.value = '2026'
        with self.assertRaisesRegex(ValueError, 'different actual years'):
            c.fetch_composition(fetcher(edit('hist03z1', relabel)))

    def test_rejects_mixed_editions(self):
        html = (FIX / FILES[c.PAGE]).read_text().replace('hist02z1_fy2027', 'hist02z1_fy2026', 1)
        with self.assertRaisesRegex(ValueError, 'single OMB'):
            c.fetch_composition(fetcher({c.PAGE: html.encode()}))

    def test_validator_rejects_bad_data(self):
        def broken(fn, pattern=None):
            d = copy.deepcopy(self.data); fn(d)
            with self.assertRaises(ValueError) if pattern is None else self.assertRaisesRegex(ValueError, pattern):
                c.validate_composition(d)
        broken(lambda d: d['history']['years'][-1]['items'][0].update(value=1))
        broken(lambda d: d['years'][3].update(total=-5))
        broken(lambda d: item(d['receipts']['years'][2], 'Excise taxes').update(value=-10, pctGdp=-10 / d['receipts']['years'][2]['gdpMillion'] * 100), 'negative')
        broken(lambda d: d['history']['years'][2]['items'][3].update(pctGdp=9.9), 'GDP shares')
        broken(lambda d: d['receipts']['years'][0].update(taxTotal=1), 'taxTotal')
        broken(lambda d: d['receipts']['years'].pop())
        broken(lambda d: d['history']['years'][1]['items'].reverse(), 'categories')
        broken(lambda d: d['edition'].update(budgetYear=2025), 'actual year')
        broken(lambda d: d['sources'][0].pop('attribution'))
        short = copy.deepcopy(self.data)
        for view in (short, short['history'], short['receipts'], short['gdp']):
            view['years'] = view['years'][-4:]
        with self.assertRaisesRegex(ValueError, 'fewer than 5'):
            c.validate_composition(short)


if __name__ == '__main__':
    unittest.main()
