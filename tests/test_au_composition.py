import copy
import io
import unittest
from pathlib import Path
import openpyxl
from pipeline.au import composition as c

FIX = Path(__file__).parent / 'fixtures/au/composition'
BASE = 'https://www.abs.gov.au/statistics/economy/government/'
FILES = {c.GFS_PAGE: 'government-finance-statistics-annual.html', c.TAX_PAGE: 'taxation-revenue-australia.html', c.GDP_URL: 'gdp.csv',
         BASE + 'government-finance-statistics-annual/2024-25/55120DO021_202425.xlsx': '55120DO021_202425.xlsx',
         BASE + 'government-finance-statistics-annual/2024-25/55120DO002_202425.xlsx': '55120DO002_202425.xlsx',
         BASE + 'taxation-revenue-australia/2024-25/55060DO001_202425.xlsx': '55060DO001_202425.xlsx',
         BASE + 'taxation-revenue-australia/2024-25/55060DO002_202425.xlsx': '55060DO002_202425.xlsx'}


def fetcher(overrides=None):
    calls = []
    def fetch(url, headers=None):
        calls.append((url, headers))
        if overrides and url in overrides:
            return overrides[url]
        return (FIX / FILES[url]).read_bytes()
    fetch.calls = calls
    return fetch


def edit(name, sheet, fn):
    book = openpyxl.load_workbook(FIX / name)
    fn(book[sheet])
    out = io.BytesIO(); book.save(out)
    return out.getvalue()


class CompositionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.fetch = fetcher()
        cls.data = c.fetch_composition(cls.fetch)

    def test_live_vintage_values(self):
        d = self.data
        self.assertEqual((d['asOf'], d['publicationDate'], d['level']), ('2024-25', '2026-04-21', 'all'))
        last = d['history']['years'][-1]
        self.assertEqual(last['total'], 1030511)
        self.assertEqual(next(i['value'] for i in last['items'] if i['name'] == 'Social protection'), 311746)
        self.assertEqual(len(d['history']['years']), 10)
        self.assertEqual(d['history']['years'][0]['year'], '2015-16')
        self.assertNotIn('debtInterest', last)
        c.validate_composition(d)

    def test_gdp_is_july_to_june_and_shares_derived(self):
        gdp = {r['year']: r['nominalMillion'] for r in self.data['gdp']['years']}
        self.assertEqual(gdp['2024-25'], 2777012)
        last = self.data['history']['years'][-1]
        self.assertAlmostEqual(last['totalPctGdp'], 1030511 / 2777012 * 100, places=3)
        self.assertIn(({'Accept': 'application/vnd.sdmx.data+csv'}), [h for u, h in self.fetch.calls if u == c.GDP_URL])

    def test_incomplete_gdp_year_is_null(self):
        lines = (FIX / 'gdp.csv').read_text().splitlines()
        raw = '\n'.join(l for l in lines if ',2025-Q1,' not in l).encode()
        data = c.fetch_composition(fetcher({c.GDP_URL: raw}))
        row = next(r for r in data['history']['years'] if r['year'] == '2024-25')
        self.assertIsNone(row['totalPctGdp']); self.assertIsNone(row['items'][0]['pctGdp'])
        c.validate_composition(data)

    def test_receipts_reconcile_to_gfs_revenue_and_tax(self):
        r = self.data['receipts']['years'][-1]
        self.assertEqual((r['total'], r['taxTotal']), (1022351, 838974))
        self.assertEqual(next(i['value'] for i in r['items'] if i['name'] == 'Goods and services tax (GST)'), 94790)
        self.assertEqual(next(i['value'] for i in r['items'] if i['name'] == c.ROUNDING), -1)
        self.assertAlmostEqual(sum(i['value'] for i in r['items']), r['total'])

    def test_commonwealth_level_is_separate_with_debt_transactions(self):
        cth = self.data['levels']['commonwealth']
        last = cth['history']['years'][-1]
        self.assertEqual((last['total'], last['debtInterest']), (745029, 42348))
        self.assertEqual(next(i['value'] for i in last['items'] if i['name'] == 'Health'), 118155)
        self.assertEqual(cth['receipts']['years'][-1]['taxTotal'], 675173)
        self.assertNotEqual(self.data['history']['years'][-1]['total'], last['total'])

    def test_parser_rejects_broken_reconciliation(self):
        def bump(ws):
            for row in ws.iter_rows():
                if str(row[0].value).strip() == 'Health':
                    row[10].value += 50
        raw = edit('55120DO021_202425.xlsx', 'Table_4', bump)
        url = BASE + 'government-finance-statistics-annual/2024-25/55120DO021_202425.xlsx'
        with self.assertRaisesRegex(ValueError, 'reconciliation'):
            c.fetch_composition(fetcher({url: raw}))

    def test_parser_rejects_tax_mismatch_and_mixed_vintages(self):
        def bump(ws):
            for row in ws.iter_rows():
                if str(row[0].value).strip() == 'Total taxation revenue':
                    row[10].value += 100
        url = BASE + 'taxation-revenue-australia/2024-25/55060DO002_202425.xlsx'
        with self.assertRaisesRegex(ValueError, 'taxation'):
            c.fetch_composition(fetcher({url: edit('55060DO002_202425.xlsx', 'Table_2', bump)}))
        html = (FIX / FILES[c.GFS_PAGE]).read_text().replace('55120DO002_202425', '55120DO002_202324', 1)
        with self.assertRaisesRegex(ValueError, 'single latest'):
            c.fetch_composition(fetcher({c.GFS_PAGE: html.encode()}))

    def test_validator_rejects_bad_data(self):
        def broken(fn):
            d = copy.deepcopy(self.data); fn(d)
            with self.assertRaises(ValueError):
                c.validate_composition(d)
        broken(lambda d: d['history']['years'][-1]['items'][0].update(value=1))
        broken(lambda d: d['years'][3].update(total=-5))
        short = copy.deepcopy(self.data)
        for view in (short, short['history'], short['receipts']):
            view['years'] = view['years'][-4:]
        with self.assertRaisesRegex(ValueError, 'fewer than 5'):
            c.validate_composition(short)
        broken(lambda d: d['levels']['commonwealth']['receipts']['years'][0].update(taxTotal=1))
        broken(lambda d: d['history']['years'][2]['items'][3].update(pctGdp=9.9))
        broken(lambda d: d['levels'].pop('commonwealth'))
        broken(lambda d: d['sources'][0].pop('attribution'))
        broken(lambda d: d['receipts']['years'].pop())


if __name__ == '__main__':
    unittest.main()
