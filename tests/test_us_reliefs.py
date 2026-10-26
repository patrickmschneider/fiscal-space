"""US Treasury tax expenditures adapter on trimmed real Treasury fixtures (FY2027 edition, public domain).

Fixtures: the landing-page link block, the workbook reduced to Tables 1, 2a, 2b and 5 (values only),
pypdf text of PDF pages 1-26 (descriptions and Table 1) and the real first PDF page.
"""
import copy
import gzip
import io
import re
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

from pipeline.us import reliefs
from pipeline.us.reliefs import fetch_reliefs, links, parse_pdf, parse_reliefs, pdf_text, validate_reliefs

FIXTURES = Path(__file__).parent / 'fixtures/us/reliefs'
LANDING = (FIXTURES / 'landing.html').read_bytes()
XLSX = (FIXTURES / 'workbook.xlsx').read_bytes()
PAGES = gzip.decompress((FIXTURES / 'pdf-pages.txt.gz').read_bytes()).decode('utf-8').split('\f')
PAGE1 = (FIXTURES / 'page1.pdf').read_bytes()
XLSX_URL = 'https://home.treasury.gov/system/files/131/Tax-Expenditures-FY2027.xlsx'
PDF_URL = 'https://home.treasury.gov/system/files/131/Tax-Expenditures-FY2027.pdf'


def fetcher(url, headers=None):
    return {reliefs.LANDING: LANDING, XLSX_URL: XLSX, PDF_URL: b'%PDF-placeholder'}[url]


def item(data, code):
    return next(r for r in data['reliefs'] if r['id'] == code)


def parse(xlsx=XLSX, pages=PAGES):
    return parse_reliefs(xlsx, pages, 2027, XLSX_URL, PDF_URL)


def edited_workbook(sheet, number, column, value, *more):
    """Workbook with cells changed; more edits follow as further (sheet, number, column, value) groups."""
    book = openpyxl.load_workbook(io.BytesIO(XLSX))
    edits = [(sheet, number, column, value)] + [tuple(more[i:i + 4]) for i in range(0, len(more), 4)]
    for sheet, number, column, value in edits:
        ws = next(book[s] for s in book.sheetnames if s.startswith(sheet + ' '))
        row = next(r for r in ws.iter_rows() if r[0].value == number)
        row[column].value = value
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


class TreasuryTaxExpenditureTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with mock.patch.object(reliefs, 'pdf_text', return_value=PAGES):
            cls.data = fetch_reliefs(fetcher)

    def test_landing_page_gives_latest_edition(self):
        self.assertEqual(links(LANDING.decode()), (2027, XLSX_URL, PDF_URL))
        newer = LANDING.decode().replace('<h3 id=prior-years>', '<a href=/system/files/131/Tax-Expenditures-FY2028.xlsx>x</a><a href=/system/files/131/Tax-Expenditures-FY2028.pdf>x</a><h3 id=prior-years>')
        self.assertEqual(links(newer)[0], 2028)
        with self.assertRaises(ValueError):
            links(LANDING.decode().replace('Tax-Expenditures-FY2027.pdf', 'Tax-Expenditures-FY2027-draft.htm'))
        with self.assertRaises(ValueError):
            links('<title>Page not found</title>')

    def test_pdf_text_reads_real_page(self):
        self.assertIn('TAX EXPENDITURES FISCAL YEAR 2027', pdf_text(PAGE1)[0])

    def test_inventory_years_and_values(self):
        d = self.data
        self.assertEqual((len(d['reliefs']), d['years'][0], d['years'][-1], d['asOf'], d['defaultYear'], d['edition']),
                         (175, '2024-25', '2034-35', '2025-12-16', '2025-26', 'FY2027'))
        health = item(d, '130')
        self.assertEqual((health['estimates'][1]['costMillion'], health['estimates'][1]['fiscalYear'], health['estimates'][1]['status'], health['taxHead'], health['functionalCategory']),
                         (295660.0, 'FY2026', 'forecast', 'Individual income tax', 'Health'))
        self.assertEqual(health['estimates'][0]['status'], 'estimate')
        self.assertIn('payroll tax receipts', health['notes'])
        cfc = item(d, '4')
        self.assertEqual((cfc['taxHead'], cfc['normalTaxMethod'], cfc['estimates'][1]['corporateMillion']), ('Corporate income tax', True, 31190.0))
        self.assertTrue(item(d, '12')['description'].startswith('This is similar to the above provision'))
        self.assertTrue(item(d, '53')['description'].endswith('remaining expenses.'))  # Function intro paragraph dropped.
        self.assertEqual(item(d, '96')['name'], 'Credit to holders of Gulf and Midwest Tax Credit Bonds')

    def test_split_outlays_and_source_issues(self):
        d = self.data
        tonnage = item(d, '86')
        self.assertEqual((tonnage['taxHead'], tonnage['splitDerived']), ('Corporate income tax', False))
        self.assertTrue(any('Tonnage tax' in i for i in d['sourceIssues']))
        eitc = item(d, '164')
        self.assertEqual((eitc['estimates'][1]['costMillion'], eitc['estimates'][1]['outlayMillion'], eitc['hasOutlayEffect']), (670.0, 68860.0, True))
        self.assertIsNone(item(d, '1')['estimates'][1]['outlayMillion'])
        rebate = item(d, '165')
        self.assertEqual(rebate['estimates'][2]['outlayMillion'], 0.0)
        self.assertEqual(d['derivedTotals']['outlays'][1]['count'], 23)

    def test_missing_split_row_is_derived_and_flagged(self):
        book = openpyxl.load_workbook(io.BytesIO(XLSX))
        ws = book['Table 2b - Individual']
        row = next(r for r in ws.iter_rows() if r[1].value == 'Deduction for car interest payments')
        for c in row:
            c.value = None
        out = io.BytesIO()
        book.save(out)
        d = parse(out.getvalue())
        car = item(d, '85')
        self.assertEqual((car['splitDerived'], car['estimates'][1]['individualMillion'], car['taxHead']), (True, 1720.0, 'Individual income tax'))

    def test_no_official_total_and_derived_sums(self):
        d = self.data
        self.assertIsNone(d['aggregate']['costMillion'])
        total = d['derivedTotals']['all'][1]
        self.assertEqual((total['year'], total['quantifiedCount'], total['quantifiedSumMillion']),
                         ('2025-26', 175, sum(r['estimates'][1]['costMillion'] for r in d['reliefs'])))
        self.assertEqual(d['retirement']['codes'], [str(n) for n in range(148, 155)])
        self.assertEqual(d['retirement']['totals'][1]['quantifiedSumMillion'], 302790.0)
        self.assertIsNone(d['isa'])
        self.assertIsNone(d['superannuation'])

    def test_group_subheadings_end_with_their_group(self):
        # Treasury Table 1: 'Exclusion of other employee benefits:' covers 155-159 only; the blind/elderly deductions,
        # casualty losses, EITC and recovery rebates (160-165) stand alone, as does 167 after the Social Security group.
        sub = {r['provisionNumber']: r['tableSubheading'] for r in self.data['reliefs']}
        self.assertEqual({sub[n] for n in range(148, 155)}, {reliefs.RETIREMENT})
        self.assertEqual({sub[n] for n in range(155, 160)}, {'Exclusion of other employee benefits'})
        self.assertEqual({sub[n] for n in range(160, 166)}, {None})
        self.assertEqual((sub[147], sub[166], sub[167], sub[168]), (None, 'Exclusion of social security benefits', None, None))
        self.assertEqual((sub[69], sub[70], sub[129]), ('Housing', 'Commerce', 'Training, employment, and social services'))

    def test_split_words_are_rejoined(self):
        words = reliefs.vocabulary(['military personnel and their qualifying children over time; investments in tangible property; a tax'])
        self.assertEqual(reliefs.rejoin('provided militar y personnel in t heir qu alifying homes over ti me', words),
                         'provided military personnel in their qualifying homes over time')
        self.assertEqual(reliefs.rejoin('investments in tangible property, a t ax, a n item', words), 'investments in tangible property, a tax, a n item')
        d = self.data
        self.assertIn('military personnel', item(d, '1')['description'])
        self.assertIn('three or more qualifying children', item(d, '164')['description'])
        self.assertFalse(any(re.search(r'\b(?:militar y|qu alifying|ti me|t o|o f|th e)\b', r['description'] or '') for r in d['reliefs']))

    def test_workbook_pdf_disagreement_fails(self):
        with self.assertRaisesRegex(ValueError, 'disagrees with PDF Table 1'):
            parse(edited_workbook('Table 1', 130, 2, 279190, 'Table 2b', 130, 2, 279190))

    def test_corporate_individual_inconsistency_fails(self):
        with self.assertRaisesRegex(ValueError, 'corporate \\+ individual'):
            parse(edited_workbook('Table 2a', 4, 2, 30000))

    def test_ten_year_total_and_text_values_checked(self):
        with self.assertRaisesRegex(ValueError, 'ten-year total'):
            parse(edited_workbook('Table 1', 8, 13, 1))
        with self.assertRaisesRegex(ValueError, 'non-numeric'):
            parse(edited_workbook('Table 1', 8, 4, 'n/a'))

    def test_pdf_layout_and_edition_checks(self):
        with self.assertRaisesRegex(ValueError, 'not the FY2028'):
            parse_pdf(PAGES, 2028)
        with self.assertRaisesRegex(ValueError, 'years changed'):
            parse_reliefs(XLSX, PAGES, 2028, XLSX_URL, PDF_URL)
        broken = [p.replace('125. ', '125) ') for p in PAGES]
        with self.assertRaises(ValueError):
            parse(pages=broken)

    def test_validator_rejects_bad_data(self):
        validate_reliefs(self.data)
        cases = [
            lambda d: d['aggregate'].update(costMillion=2_000_000),
            lambda d: d.update(isa={'history': []}),
            lambda d: d['reliefs'].append(copy.deepcopy(d['reliefs'][0])),
            lambda d: d['derivedTotals']['all'][1].update(quantifiedSumMillion=1.0),
            lambda d: d.update(reliefs=d['reliefs'][:100]),
            lambda d: item(d, '130')['estimates'][1].update(costMillion=None),
            lambda d: item(d, '130')['estimates'][1].update(costMillion=2_957_000.0, individualMillion=2_957_000.0),
            lambda d: item(d, '4')['estimates'][1].update(corporateMillion=1.0),
            lambda d: item(d, '4')['estimates'][0].update(status='forecast'),
            lambda d: d['years'].pop(),
        ]
        for mutate in cases:
            bad = copy.deepcopy(self.data)
            mutate(bad)
            with self.assertRaises(ValueError):
                validate_reliefs(bad)


if __name__ == '__main__':
    unittest.main()
